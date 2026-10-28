import io
import csv
import uuid
import datetime as dt
from typing import List, Dict, Any, Optional, Tuple
from sqlalchemy import select, func
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.models.academic import CourseSection, Course, Enrollment, Department
from app.models.attendance import AttendanceSession, AttendanceRecord, AttendancePolicy
from app.models.user import Student, User, Teacher


class ReportService:
    @staticmethod
    async def get_section_attendance_matrix(
        db: AsyncSession,
        section_id: uuid.UUID,
    ) -> Dict[str, Any]:
        """
        Build full attendance matrix for a section:
        All completed sessions as columns, all enrolled students as rows, status in cells, percentages calculated.
        """
        # Load section and course
        sec_stmt = (
            select(CourseSection)
            .options(
                selectinload(CourseSection.course).selectinload(Course.department),
                selectinload(CourseSection.enrollments).selectinload(Enrollment.student).selectinload(Student.user),
                selectinload(CourseSection.attendance_sessions),
            )
            .where(CourseSection.id == section_id)
        )
        res = await db.execute(sec_stmt)
        section = res.scalars().first()
        if not section:
            return {}

        # Completed sessions sorted by date
        sessions = sorted(
            [s for s in section.attendance_sessions if s.state == "COMPLETED"],
            key=lambda x: (x.session_date, x.start_time),
        )

        session_ids = [s.id for s in sessions]

        # Load all records for these sessions
        records_stmt = (
            select(AttendanceRecord)
            .where(AttendanceRecord.session_id.in_(session_ids))
        ) if session_ids else None
        
        records_map: Dict[Tuple[uuid.UUID, uuid.UUID], str] = {}
        if records_stmt is not None:
            records = (await db.execute(records_stmt)).scalars().all()
            for r in records:
                records_map[(r.student_id, r.session_id)] = r.status

        students_data = []
        total_sessions = len(sessions)

        for enrollment in sorted(section.enrollments, key=lambda e: e.student.student_id_number):
            student = enrollment.student
            user = student.user

            session_statuses = []
            present_cnt = 0
            late_cnt = 0
            absent_cnt = 0
            excused_cnt = 0

            for sess in sessions:
                st = records_map.get((student.id, sess.id), "ABSENT")
                session_statuses.append(st)
                if st == "PRESENT":
                    present_cnt += 1
                elif st == "LATE":
                    late_cnt += 1
                elif st == "EXCUSED":
                    excused_cnt += 1
                else:
                    absent_cnt += 1

            if total_sessions > 0:
                pct = round(((present_cnt + (late_cnt * 0.5) + excused_cnt) / total_sessions) * 100.0, 1)
            else:
                pct = 100.0

            students_data.append({
                "student_id": str(student.id),
                "student_id_number": student.student_id_number,
                "student_name": user.full_name,
                "present": present_cnt,
                "late": late_cnt,
                "absent": absent_cnt,
                "excused": excused_cnt,
                "percentage": pct,
                "session_statuses": session_statuses,
            })

        return {
            "section_id": str(section.id),
            "course_code": section.course.code,
            "course_name": section.course.name,
            "section_name": section.section_name,
            "department_name": section.course.department.name,
            "total_sessions": total_sessions,
            "sessions": [
                {
                    "id": str(s.id),
                    "title": s.title,
                    "date": s.session_date.strftime("%Y-%m-%d"),
                }
                for s in sessions
            ],
            "students": students_data,
        }

    @staticmethod
    def generate_csv(matrix: Dict[str, Any]) -> str:
        """Export attendance matrix to CSV."""
        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow(["AttendX - Attendance Report"])
        writer.writerow(["Course:", f"{matrix.get('course_code')} - {matrix.get('course_name')}"])
        writer.writerow(["Section:", matrix.get("section_name")])
        writer.writerow(["Generated At:", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        writer.writerow([])

        # Header
        headers = ["Student ID", "Full Name", "Present", "Late", "Absent", "Excused", "Attendance %"]
        for sess in matrix.get("sessions", []):
            headers.append(f"{sess['date']} ({sess['title'][:10]})")
        writer.writerow(headers)

        # Rows
        for s in matrix.get("students", []):
            row = [
                s["student_id_number"],
                s["student_name"],
                s["present"],
                s["late"],
                s["absent"],
                s["excused"],
                f"{s['percentage']}%",
            ]
            row.extend(s.get("session_statuses", []))
            writer.writerow(row)

        return output.getvalue()

    @staticmethod
    def generate_excel(matrix: Dict[str, Any]) -> bytes:
        """Export styled Excel sheet with conditional formatting."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Summary"

        # Brand header
        header_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        ws.cell(row=1, column=1, value=f"AttendX - {matrix.get('course_code')} {matrix.get('course_name')} ({matrix.get('section_name')})").font = header_font
        ws.cell(row=2, column=1, value=f"Department: {matrix.get('department_name')} | Total Sessions: {matrix.get('total_sessions')}").font = Font(size=10, italic=True)

        cols = ["Student ID", "Full Name", "Present", "Late", "Absent", "Excused", "Percentage"]
        for sess in matrix.get("sessions", []):
            cols.append(f"{sess['date']}")

        # Header row styling
        hdr_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        hdr_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

        for col_idx, text in enumerate(cols, 1):
            cell = ws.cell(row=4, column=col_idx, value=text)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center")

        # Student rows
        row_idx = 5
        for s in matrix.get("students", []):
            ws.cell(row=row_idx, column=1, value=s["student_id_number"])
            ws.cell(row=row_idx, column=2, value=s["student_name"])
            ws.cell(row=row_idx, column=3, value=s["present"])
            ws.cell(row=row_idx, column=4, value=s["late"])
            ws.cell(row=row_idx, column=5, value=s["absent"])
            ws.cell(row=row_idx, column=6, value=s["excused"])
            pct_cell = ws.cell(row=row_idx, column=7, value=f"{s['percentage']}%")

            # Highlight low attendance
            if s["percentage"] < 75.0:
                pct_cell.font = Font(color="DC2626", bold=True)

            for sess_idx, status_val in enumerate(s.get("session_statuses", []), 8):
                c = ws.cell(row=row_idx, column=sess_idx, value=status_val)
                if status_val == "PRESENT":
                    c.font = Font(color="16A34A")
                elif status_val == "LATE":
                    c.font = Font(color="D97706")
                elif status_val == "ABSENT":
                    c.font = Font(color="DC2626")

            row_idx += 1

        # Auto-adjust column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def generate_pdf(matrix: Dict[str, Any]) -> bytes:
        """Export professional PDF summary using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "Title",
            parent=styles["Heading1"],
            fontSize=18,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=6,
        )
        subtitle_style = ParagraphStyle(
            "Subtitle",
            parent=styles["Normal"],
            fontSize=11,
            textColor=colors.HexColor("#475569"),
            spaceAfter=15,
        )

        elements.append(Paragraph(f"AttendX Attendance Report — {matrix.get('course_code')}", title_style))
        elements.append(Paragraph(
            f"<b>Course:</b> {matrix.get('course_name')} | <b>Section:</b> {matrix.get('section_name')} | <b>Total Classes:</b> {matrix.get('total_sessions')}",
            subtitle_style,
        ))

        # Table data
        table_data = [["Student ID", "Full Name", "Present", "Late", "Absent", "Excused", "Overall %"]]
        for s in matrix.get("students", []):
            table_data.append([
                s["student_id_number"],
                s["student_name"][:20],
                str(s["present"]),
                str(s["late"]),
                str(s["absent"]),
                str(s["excused"]),
                f"{s['percentage']}%",
            ])

        t = Table(table_data, colWidths=[90, 160, 50, 45, 50, 55, 70])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("TOPPADDING", (0, 0), (-1, 0), 6),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ALIGN", (2, 1), (-1, -1), "CENTER"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
        ]))
        elements.append(t)

        doc.build(elements)
        return buffer.getvalue()
